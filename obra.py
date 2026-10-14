"""Datos reproducibles de obra y gráficos obtenidos exclusivamente del Excel."""
from pathlib import Path
import argparse
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import PercentFormatter
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
EXCEL = ROOT / "datos_obra.xlsx"
OUT = ROOT / "graficos"


def generar_excel():
    # Edificio residencial de cinco pisos: cantidades y precios unitarios en soles.
    actividades = [
        ("P01", "Obras provisionales", "glb", 1, 85000, 1, 3, 0),
        ("P02", "Movimiento de tierras", "m³", 1800, 42, 2, 5, 1),
        ("P03", "Cimentaciones", "m³", 420, 980, 4, 9, 1),
        ("P04", "Estructuras de concreto", "m³", 1150, 1250, 8, 22, 2),
        ("P05", "Albañilería", "m²", 4200, 88, 15, 27, 2),
        ("P06", "Instalaciones sanitarias", "pto", 260, 620, 17, 30, 1),
        ("P07", "Instalaciones eléctricas", "pto", 420, 390, 17, 31, 2),
        ("P08", "Tarrajeos", "m²", 8500, 32, 22, 33, 2),
        ("P09", "Pisos y revestimientos", "m²", 3600, 115, 27, 37, 1),
        ("P10", "Carpintería y vidrios", "m²", 780, 340, 30, 38, 2),
        ("P11", "Pintura", "m²", 9200, 19, 33, 39, 1),
        ("P12", "Pruebas y entrega", "glb", 1, 65000, 38, 40, 1),
    ]
    cols = ["codigo", "partida", "unidad", "cantidad_total", "precio_unitario", "inicio", "fin", "retraso_semanas"]
    partidas = pd.DataFrame(actividades, columns=cols)
    partidas["presupuesto"] = partidas.cantidad_total * partidas.precio_unitario
    partidas["peso"] = partidas.presupuesto / partidas.presupuesto.sum()
    rng = np.random.default_rng(19)
    registros = []
    inicio = pd.Timestamp("2026-01-05")
    corte = 28
    for p in partidas.itertuples():
        semanas = np.arange(p.inicio, p.fin + 1)
        perfil = np.sin(np.linspace(0.15, np.pi - 0.15, len(semanas)))
        plan = perfil / perfil.sum() * p.cantidad_total
        real = plan * rng.uniform(0.88, 1.08, len(plan))
        # Conserva el metrado total; cambia ritmo y desplaza por retraso.
        real *= p.cantidad_total / real.sum()
        for semana in range(1, 41):
            ip = semana - p.inicio
            ir = semana - p.inicio - p.retraso_semanas
            qp = float(plan[ip]) if 0 <= ip < len(plan) else 0.0
            qr = (float(real[ir]) if 0 <= ir < len(real) else 0.0) if semana <= corte else np.nan
            registros.append((semana, inicio + pd.Timedelta(weeks=semana-1), p.codigo, qp, qr))
    datos = pd.DataFrame(registros, columns=["semana", "fecha", "codigo", "cantidad_plan", "cantidad_real"])
    parametros = pd.DataFrame({"parametro": ["Proyecto", "Moneda", "Semana de corte", "Semanas del plan", "Fuente"],
                               "valor": ["Edificio residencial de 5 pisos", "PEN (S/)", corte, 40, "Datos sintéticos; semilla 19"]})
    with pd.ExcelWriter(EXCEL, engine="openpyxl") as writer:
        parametros.to_excel(writer, sheet_name="Parametros", index=False)
        partidas.to_excel(writer, sheet_name="Partidas", index=False)
        datos.to_excel(writer, sheet_name="Avance_semanal", index=False)
        for ws in writer.book:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="6535A5")
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="center")
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0.00"
            for col in ws.columns:
                width = min(52, max(len(str(c.value or "")) for c in col) + 3)
                ws.column_dimensions[get_column_letter(col[0].column)].width = width
        writer.book["Avance_semanal"].column_dimensions["B"].width = 16
        for cell in writer.book["Avance_semanal"]["B"][1:]:
            cell.number_format = "dd/mm/yyyy"
        for cell in writer.book["Partidas"]["J"][1:]:
            cell.number_format = "0.0%"


def leer_y_calcular():
    partidas = pd.read_excel(EXCEL, sheet_name="Partidas")
    datos = pd.read_excel(EXCEL, sheet_name="Avance_semanal")
    parametros = pd.read_excel(EXCEL, sheet_name="Parametros").set_index("parametro").valor
    corte = int(parametros["Semana de corte"])
    assert partidas.codigo.is_unique, "Códigos de partida duplicados"
    assert not datos.duplicated(["semana", "codigo"]).any(), "Registros semanales duplicados"
    assert set(datos.codigo) == set(partidas.codigo), "Partidas sin correspondencia"
    assert 1 <= corte <= 40, "Corte fuera del cronograma"
    assert (partidas.cantidad_total > 0).all() and (partidas.precio_unitario > 0).all()
    assert datos.loc[datos.semana <= corte, "cantidad_real"].notna().all(), "Faltan datos al corte"
    assert (datos.cantidad_plan >= 0).all() and (datos.cantidad_real.dropna() >= 0).all()
    presupuesto = (partidas.cantidad_total * partidas.precio_unitario).sum()
    detalle = datos.merge(partidas, on="codigo", validate="many_to_one")
    detalle["valor_plan"] = detalle.cantidad_plan * detalle.precio_unitario
    detalle["valor_real"] = detalle.cantidad_real * detalle.precio_unitario
    semanal = detalle.groupby("semana")[["valor_plan", "valor_real"]].sum(min_count=1)
    semanal["plan_acum"] = semanal.valor_plan.cumsum() / presupuesto
    semanal["real_acum"] = semanal.valor_real.cumsum() / presupuesto
    semanal.loc[semanal.index > corte, ["valor_real", "real_acum"]] = np.nan
    avance = detalle[detalle.semana <= corte].groupby("codigo")[["cantidad_plan", "cantidad_real"]].sum()
    avance = partidas.set_index("codigo").join(avance)
    avance["plan_pct"] = avance.cantidad_plan / avance.cantidad_total
    avance["real_pct"] = avance.cantidad_real / avance.cantidad_total
    assert np.isclose(semanal.plan_acum.iloc[-1], 1), "El plan no suma 100 %"
    assert (avance.real_pct <= 1 + 1e-8).all(), "Avance superior al metrado"
    return semanal, avance, corte, presupuesto


def graficar():
    semanal, avance, corte, presupuesto = leer_y_calcular()
    OUT.mkdir(exist_ok=True)
    plt.rcParams.update({"figure.facecolor": "#09080D", "axes.facecolor": "#09080D",
        "text.color": "#F1EAFB", "axes.labelcolor": "#DDD0EF", "xtick.color": "#DDD0EF",
        "ytick.color": "#DDD0EF", "axes.edgecolor": "#554064", "grid.color": "#33253E",
        "font.size": 11, "savefig.facecolor": "#09080D"})
    plan, real = "#9163D9", "#E0B5FF"
    fig, ax = plt.subplots(figsize=(12, 6.5), layout="constrained")
    x = np.r_[0, semanal.index]
    ax.plot(x, np.r_[0, semanal.plan_acum], color=plan, lw=3, label="Planificado")
    ax.plot(np.r_[0, semanal.index[:corte]], np.r_[0, semanal.real_acum.iloc[:corte]], color=real, lw=3, label="Ejecutado al corte")
    ax.fill_between(semanal.index[:corte], semanal.real_acum.iloc[:corte], semanal.plan_acum.iloc[:corte], color=plan, alpha=0.18)
    ax.axvline(corte, color=real, ls=":", alpha=0.7)
    p, r = semanal.loc[corte, ["plan_acum", "real_acum"]]
    ax.set_title(f"CURVA S · Edificio residencial de 5 pisos\nCorte: semana {corte} | Plan {p:.1%} · Real {r:.1%} · Brecha {(r-p)*100:+.1f} pp", loc="left", pad=20)
    ax.set(xlabel="Semana de obra", ylabel="Avance acumulado ponderado por presupuesto", xlim=(0, 40), ylim=(0, 1.05))
    ax.yaxis.set_major_formatter(PercentFormatter(1))
    ax.grid(alpha=0.6)
    ax.legend(facecolor="#191020", labelcolor="#F1EAFB", loc="upper left")
    fig.savefig(OUT / "curva_s.png", dpi=180)
    plt.close(fig)
    fig, ax = plt.subplots(figsize=(12, 8), layout="constrained")
    y = np.arange(len(avance))
    ax.barh(y-0.19, avance.plan_pct, height=0.34, color=plan, label="Planificado al corte")
    ax.barh(y+0.19, avance.real_pct, height=0.34, color=real, label="Ejecutado al corte")
    for i, row in enumerate(avance.itertuples()):
        ax.text(row.plan_pct+0.012, i-0.19, f"{row.plan_pct:.0%}", va="center", fontsize=9, color=plan)
        ax.text(row.real_pct+0.012, i+0.19, f"{row.real_pct:.0%}", va="center", fontsize=9, color=real)
    ax.set_yticks(y, avance.partida)
    ax.invert_yaxis()
    ax.set(xlim=(0, 1.15), xlabel="Cantidad acumulada / cantidad total de la partida")
    ax.xaxis.set_major_formatter(PercentFormatter(1))
    ax.set_title(f"AVANCE POR PARTIDAS · Semana {corte}\nPresupuesto base: S/ {presupuesto:,.0f} · Datos sintéticos", loc="left", pad=20)
    ax.grid(axis="x", alpha=0.5)
    ax.set_axisbelow(True)
    ax.legend(facecolor="#191020", labelcolor="#F1EAFB", loc="lower right")
    fig.savefig(OUT / "avance_partidas.png", dpi=180)
    plt.close(fig)
    semanal.to_csv(OUT / "resumen_semanal.csv", encoding="utf-8-sig")
    print(f"Excel: {EXCEL}\nPresupuesto: S/ {presupuesto:,.2f}\nCorte {corte}: plan {p:.2%}, real {r:.2%}\nGráficos: {OUT}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--generar-datos", action="store_true", help="Crea o reemplaza el Excel sintético")
    args = parser.parse_args()
    if args.generar_datos or not EXCEL.exists():
        generar_excel()
    graficar()
